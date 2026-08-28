"""Enactus_Evidence_Tracker.xlsx — the year-round documentation workbook the
Impact & Financial Reporting Review expects, with a Summary sheet that
computes the Standardized Impact Page numbers from the raw rows."""
import sys
sys.path = [p for p in sys.path if 'dist-packages' not in p]
sys.path.insert(0, '/tmp/claude-0/-home-user-Chores/3f438fd9-0f01-509a-9b38-4373bde88bbd/scratchpad/pylibs')

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from datetime import date

NAVY, PURPLE, LAV, YELL, GRAY = "11256A", "800080", "EFE9FF", "FFF2CC", "8A8A8A"
H = Font(name="Arial", bold=True, color="FFFFFF", size=10)
B = Font(name="Arial", size=10)
EX = Font(name="Arial", size=10, italic=True, color=GRAY)
HFILL = PatternFill("solid", fgColor=NAVY)
YFILL = PatternFill("solid", fgColor=YELL)
LFILL = PatternFill("solid", fgColor=LAV)
thin = Side(style="thin", color="C9C2DA")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)

wb = Workbook()

def sheet(name, headers, widths):
    ws = wb.create_sheet(name)
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.border = H, HFILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    return ws

# ---------------- README ----------------
rd = wb.active
rd.title = "README"
rd.column_dimensions["A"].width = 30
rd.column_dimensions["B"].width = 90
rows = [
    ("Enactus Evidence Tracker — Fliq / Enactus HKUST", ""),
    ("", ""),
    ("Impact window (edit these two cells — everything filters on them)", ""),
    ("Window start", date(2026, 9, 1)),
    ("Window end", date(2027, 8, 31)),
    ("", ""),
    ("How to use", "Log every beneficiary, session, dollar and mention in the sheets as it happens — not at competition time. The Summary sheet computes the Standardized Impact Page numbers from these rows automatically; only rows inside the impact window count."),
    ("Yellow cells", "The only cells you edit on this sheet (the window dates)."),
    ("Gray italic rows", "Example rows showing the expected format — replace them with real data."),
    ("Direct impact rule", "A beneficiary counts as Direct Impact only if Type = Direct AND both Pre and Post scores are recorded (measured change, per Enactus definitions)."),
    ("Reach rule", "Downloads, waitlist, media and event exposure are Reach — never report them as impact."),
    ("Income rule", "Only Sales / Grant / Sponsorship / Donation with Direction = Income count as Income/Revenue. Log equity investment with Direction = Excluded so it can never leak into the table."),
    ("Review readiness", "Keep receipt/consent/evidence links filled: the Impact & Financial Reporting Review may contact beneficiaries and ask for documents. Numbers freeze one month before the World Cup."),
]
for r, (a, b) in enumerate(rows, 1):
    ca, cb = rd.cell(row=r, column=1, value=a), rd.cell(row=r, column=2, value=b)
    ca.font = Font(name="Arial", bold=True, size=12 if r == 1 else 10)
    cb.font = Font(name="Arial", size=10)
    cb.alignment = Alignment(wrap_text=True, vertical="top")
for r in (4, 5):
    c = rd.cell(row=r, column=2); c.fill = YFILL; c.border = BOX
    c.number_format = "DD MMM YYYY"

# ---------------- Beneficiaries ----------------
bn = sheet("Beneficiaries",
    ["ID", "Name", "Project", "Type", "School / Channel", "Consent (Y/N)", "Contact",
     "Date engaged", "Pre score", "Post score", "Improvement %", "Measured direct?", "In window?", "Notes"],
    [7, 20, 16, 11, 22, 12, 22, 13, 10, 10, 13, 14, 11, 28])
ex = ["B-001", "Chan Tai Man", "Decode: Veritas", "Direct", "SKH Secondary School", "Y",
      "teacher@school.edu.hk", date(2026, 10, 14), 55, 82, None, None, None,
      "EXAMPLE ROW — replace. Pre/post = media-literacy test scores."]
for i, v in enumerate(ex, 1):
    c = bn.cell(row=2, column=i, value=v); c.font = EX; c.border = BOX
    if i == 8: c.number_format = "DD MMM YYYY"
for r in range(2, 202):
    bn.cell(row=r, column=11, value=f'=IF(AND(ISNUMBER(I{r}),ISNUMBER(J{r}),I{r}>0),(J{r}-I{r})/I{r},"")').number_format = "0.0%"
    bn.cell(row=r, column=12, value=f'=IF(A{r}="","",IF(AND(D{r}="Direct",ISNUMBER(I{r}),ISNUMBER(J{r})),"Yes","No"))')
    bn.cell(row=r, column=13, value=f'=IF(H{r}="","",IF(AND(H{r}>=README!$B$4,H{r}<=README!$B$5),"Yes","No"))')
    for col in (11, 12, 13):
        bn.cell(row=r, column=col).font = B

# ---------------- Sessions ----------------
sn = sheet("Sessions",
    ["Date", "Project", "School / Venue", "Session topic", "Attendees", "Facilitator (team member)", "Evidence link (register/photos)", "In window?"],
    [13, 16, 24, 26, 11, 22, 30, 11])
exs = [date(2026, 10, 14), "Decode: Veritas", "SKH Secondary School", "Claims vs evidence workshop", 32,
       "Ling Bo Zeng", "drive.google.com/…", None]
for i, v in enumerate(exs, 1):
    c = sn.cell(row=2, column=i, value=v); c.font = EX; c.border = BOX
    if i == 1: c.number_format = "DD MMM YYYY"
for r in range(2, 202):
    sn.cell(row=r, column=8, value=f'=IF(A{r}="","",IF(AND(A{r}>=README!$B$4,A{r}<=README!$B$5),"Yes","No"))').font = B

# ---------------- Revenue ----------------
rv = sheet("Revenue",
    ["Date", "Project", "Direction", "Type", "Description", "Amount (USD)", "Receipt / doc ref", "In window?"],
    [13, 16, 12, 20, 34, 13, 24, 11])
exr = [
    [date(2026, 11, 2), "Fliq", "Income", "Sales", "EXAMPLE — App subscriptions, November", 120, "stripe-2026-11.pdf", None],
    [date(2026, 11, 20), "Decode: Veritas", "Income", "Sales", "EXAMPLE — School programme fee", 250, "invoice-DV-003.pdf", None],
    [date(2026, 12, 1), "Fliq", "Excluded", "Investment (excluded)", "EXAMPLE — Angel round tranche (never in the table)", 10000, "sha-2026.pdf", None],
    [date(2026, 11, 30), "Fliq", "Expense", "Infrastructure", "EXAMPLE — Cloud + licensed sourcing", 310, "aws-2026-11.pdf", None],
]
for j, row in enumerate(exr, 2):
    for i, v in enumerate(row, 1):
        c = rv.cell(row=j, column=i, value=v); c.font = EX; c.border = BOX
        if i == 1: c.number_format = "DD MMM YYYY"
        if i == 6: c.number_format = "$#,##0"
for r in range(2, 202):
    rv.cell(row=r, column=8, value=f'=IF(A{r}="","",IF(AND(A{r}>=README!$B$4,A{r}<=README!$B$5),"Yes","No"))').font = B

# ---------------- Partners ----------------
pt = sheet("Partners",
    ["Partner", "Role in project", "What the partner provided", "What the team itself did", "Contact", "Agreement / doc ref"],
    [22, 18, 34, 34, 24, 20])
exp = ["SKH Secondary School", "Host school", "Classroom access, student cohort",
       "EXAMPLE — All curriculum design, teaching and testing by team", "vice-principal@school.edu.hk", "mou-skh.pdf"]
for i, v in enumerate(exp, 1):
    c = pt.cell(row=2, column=i, value=v); c.font = EX; c.border = BOX

# ---------------- ReachLog ----------------
rl = sheet("ReachLog",
    ["Date", "Project", "Channel", "Count (people)", "Evidence link", "In window?"],
    [13, 16, 26, 14, 32, 11])
exl = [date(2026, 11, 15), "Fliq", "App downloads (cumulative delta)", 480, "appstoreconnect screenshot", None]
for i, v in enumerate(exl, 1):
    c = rl.cell(row=2, column=i, value=v); c.font = EX; c.border = BOX
    if i == 1: c.number_format = "DD MMM YYYY"
for r in range(2, 202):
    rl.cell(row=r, column=6, value=f'=IF(A{r}="","",IF(AND(A{r}>=README!$B$4,A{r}<=README!$B$5),"Yes","No"))').font = B

# ---------------- Summary ----------------
sm = wb.create_sheet("Summary")
widths = [18, 15, 15, 12, 15, 14, 14, 16]
heads = ["Project", "Direct Impact (people)", "Indirect Impact (people)", "Reach (people)",
         "Income/Revenue (USD)", "Expenses (USD)", "Profit/Surplus (USD)", "Avg score improvement"]
for i, (h, w) in enumerate(zip(heads, widths), 1):
    c = sm.cell(row=2, column=i, value=h)
    c.font, c.fill, c.border = H, HFILL, BOX
    c.alignment = Alignment(wrap_text=True, vertical="center")
    sm.column_dimensions[get_column_letter(i)].width = w
sm.cell(row=1, column=1, value="Standardized Impact Page — computed from the raw sheets (copy these numbers to the report only after deleting example rows)").font = Font(name="Arial", bold=True, size=11)
projects = ["Decode: Veritas", "Fliq", "Other"]
for j, prj in enumerate(projects, 3):
    sm.cell(row=j, column=1, value=prj).font = Font(name="Arial", bold=True, size=10)
    sm.cell(row=j, column=2, value=f'=COUNTIFS(Beneficiaries!C:C,$A{j},Beneficiaries!L:L,"Yes",Beneficiaries!M:M,"Yes")')
    sm.cell(row=j, column=3, value=f'=COUNTIFS(Beneficiaries!C:C,$A{j},Beneficiaries!D:D,"Indirect",Beneficiaries!M:M,"Yes")')
    sm.cell(row=j, column=4, value=f'=SUMIFS(ReachLog!D:D,ReachLog!B:B,$A{j},ReachLog!F:F,"Yes")')
    sm.cell(row=j, column=5, value=f'=SUMIFS(Revenue!F:F,Revenue!B:B,$A{j},Revenue!C:C,"Income",Revenue!H:H,"Yes")')
    sm.cell(row=j, column=6, value=f'=SUMIFS(Revenue!F:F,Revenue!B:B,$A{j},Revenue!C:C,"Expense",Revenue!H:H,"Yes")')
    sm.cell(row=j, column=7, value=f'=E{j}-F{j}')
    sm.cell(row=j, column=8, value=f'=IFERROR(AVERAGEIFS(Beneficiaries!K:K,Beneficiaries!C:C,$A{j},Beneficiaries!L:L,"Yes",Beneficiaries!M:M,"Yes"),"")')
sm.cell(row=6, column=1, value="Total").font = Font(name="Arial", bold=True, size=10)
for col in range(2, 8):
    L = get_column_letter(col)
    sm.cell(row=6, column=col, value=f'=SUM({L}3:{L}5)')
for r in range(3, 7):
    for col in range(1, 9):
        c = sm.cell(row=r, column=col)
        c.border = BOX
        if c.font is None or not c.font.bold: c.font = B
        if col in (5, 6, 7): c.number_format = "$#,##0"
        if col == 8: c.number_format = "0.0%"
    if r == 6:
        for col in range(1, 9): sm.cell(row=r, column=col).fill = LFILL
sm.cell(row=8, column=1, value="Note: 'Other' rows come from any additional Enactus HKUST projects logged with Project = Other. Copy the totals into the Standardized Impact Page; the two must always match.").font = Font(name="Arial", italic=True, size=9, color=GRAY)

# ---------------- data validation ----------------
dv_prj = DataValidation(type="list", formula1='"Decode: Veritas,Fliq,Other"', allow_blank=True)
dv_typ = DataValidation(type="list", formula1='"Direct,Indirect"', allow_blank=True)
dv_yn  = DataValidation(type="list", formula1='"Y,N"', allow_blank=True)
dv_dir = DataValidation(type="list", formula1='"Income,Expense,Excluded"', allow_blank=True)
bn.add_data_validation(dv_prj); dv_prj.add("C2:C201")
bn.add_data_validation(dv_typ); dv_typ.add("D2:D201")
bn.add_data_validation(dv_yn);  dv_yn.add("F2:F201")
dv_prj2 = DataValidation(type="list", formula1='"Decode: Veritas,Fliq,Other"', allow_blank=True)
sn.add_data_validation(dv_prj2); dv_prj2.add("B2:B201")
dv_prj3 = DataValidation(type="list", formula1='"Decode: Veritas,Fliq,Other"', allow_blank=True)
dv_dir2 = DataValidation(type="list", formula1='"Income,Expense,Excluded"', allow_blank=True)
rv.add_data_validation(dv_prj3); dv_prj3.add("B2:B201")
rv.add_data_validation(dv_dir2); dv_dir2.add("C2:C201")
dv_prj4 = DataValidation(type="list", formula1='"Decode: Veritas,Fliq,Other"', allow_blank=True)
rl.add_data_validation(dv_prj4); dv_prj4.add("B2:B201")

wb.save("/home/user/Chores/enactus/Enactus_Evidence_Tracker.xlsx")
print("Wrote Enactus_Evidence_Tracker.xlsx")
